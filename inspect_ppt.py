import os
from pptx import Presentation

def inspect(path, out_path):
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return
    try:
        prs = Presentation(path)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(f"Number of slides: {len(prs.slides)}\n")
            for i, slide in enumerate(prs.slides):
                f.write(f"--- Slide {i+1} ---\n")
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        f.write(shape.text + "\n")
    except Exception as e:
        print(f"Error reading PPT: {e}")

if __name__ == "__main__":
    inspect(r"C:\Users\Dell\Downloads\图像压缩.pptx", r"e:\程序\java\MCBG\ppt_content.txt")
