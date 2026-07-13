import os
from pptx import Presentation

def inspect_ppt(path):
    if not os.path.exists(path):
        print(f"Error: File not found at {path}")
        return
        
    prs = Presentation(path)
    print(f"PPT loaded successfully. Total slides: {len(prs.slides)}")
    for idx, slide in enumerate(prs.slides):
        title = "No Title"
        # 寻找幻灯片的标题占位符
        if slide.shapes.title:
            title = slide.shapes.title.text
        else:
            # 如果没有标准标题，寻找第一个有文字的文本框作为临时标题
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    title = shape.text_frame.text.split('\n')[0][:30]
                    break
        print(f"Slide {idx}: {title}")

if __name__ == '__main__':
    path = r"C:\Users\Dell\Downloads\图像压缩.pptx"
    inspect_ppt(path)
