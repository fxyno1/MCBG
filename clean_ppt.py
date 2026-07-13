import os
import re
from pptx import Presentation

# 判定段落是否为代码段
def is_code_paragraph(text):
    code_keywords = [
        'int length', 'void Compress', 'void Traceback', 'void Output',
        'while', 'return', 'cout', 'include', 'using namespace', 'int k = 1',
        'int Lmax', 'b[i] = length', 'bmax = b[i-j+1]'
    ]
    text_stripped = text.strip()
    if any(kw in text_stripped for kw in code_keywords):
        return True
    if text_stripped.startswith('//'):
        return True
    if text_stripped.endswith('{') or text_stripped.endswith('}') or text_stripped.endswith('};'):
        return True
    if ';' in text_stripped and ('int ' in text_stripped or 'for' in text_stripped or 'while' in text_stripped or 'cout' in text_stripped):
        return True
    return False

# 规则库：匹配特定 LaTeX/AI 不自然段落，翻译为顺畅的中文自然语言
RULES = [
    # Slide 2 & general intro
    (r"计算机中常用像素点灰度值序列.*表示图像，其中.*", 
     "计算机中常用像素值序列 {p1, p2, ..., pn} 来表示图像，其中每个像素值 pi 都是 0 到 255 之间的整数。"),
    
    (r"为了节省空间，将序列分割成.*个连续段.*对于每个像素段.*", 
     "为了节省空间，需要将像素序列分割成若干个连续的段。对于每一个分段："),
    
    (r".*段长度\s*\$l\[i\]\$.*存储长度需要.*", 
     "• 段长度不超过 256 像素（记录长度需要 8 位二进制）。"),
    
    (r".*段位宽\s*\$b\[i\]\$.*存储位宽需要.*", 
     "• 段内像素的最大位宽不超过 8 位（记录位宽需要 3 位二进制）。"),
    
    (r".*每段的头部开销.*Header.*固定值.*", 
     "• 每一段的段头信息开销为固定的 11 位（8位表示长度 + 3位表示位宽）。"),
    
    (r"\$\$\s*\\min\s+\\sum_\{i=1\}\^\{m\}\s*\\left\(\s*l\[i\]\s*\\cdot\s*b\[i\]\s*\+\s*11\s*\\right\)\s*\$\$", 
     "目标是最小化存储总空间，即：所有段的（段长度 × 段位宽 + 11位段头开销）之和的最小值。"),
     
    # Slide 3
    (r"设.*的最优分段为.*最后一段为.*", 
     "设像素序列的最优分段中，最后一段所包含的像素为从第 n-L+1 到第 n 个像素（其中 L 为最后一段的长度）。"),
     
    (r"如果我们将最后一段切除，剩余的像素序列.*的分段方式.*必须是该前缀序列的最优分段.*", 
     "如果我们将最后这一段切除，那么剩余的前缀像素序列的分段方式，也必须是该前缀序列的最优分段方案。"),
     
    (r"若不然，假设该前缀序列存在一个更优分段方式，其开销为.*", 
     "否则，假设该前缀序列存在一种更省空间的划分方式，其存储开销比当前方案更小。"),
     
    (r"则用该分段与最后一段.*组合，可得到总开销为.*该开销小于原最优解开销，产生矛盾.*", 
     "那我们将这个更省空间的分段与最后一段合并，就会得到一个总空间开销更小的新方案。这与我们假设原方案已是最优相矛盾。"),
     
    (r"因此，图像压缩原问题的最优解包含了子问题的最优解，问题满足最优子结构性质.*", 
     "因此，图像压缩问题的最优解包含了子问题的最优解，满足最优子结构性质。"),
     
    # Slide 4
    (r"设\s+\$s\[i\]\$.*为前缀像素序列.*", 
     "定义 s[i] 为前 i 个像素（即 p1 到 pi）进行最优分段所需的最小存储位数。"),
     
    (r"假设当前子问题最优解的最后一段长度为.*", 
     "假设当前子问题最优解中，最后一段的长度为 k（其中 k 大于等于 1，且不超过像素数 i 和 256 的最小值）："),
     
    (r".*该段像素的最大位宽.*b_\{?\\text\{max\}\}?.*", 
     "• 该段中像素的最大位宽：为该区间内所有像素所需二进制位数的最大值。"),
     
    (r"其中\s+\$\\text\{length\}.*", 
     "其中每个像素 p 所需的二进制位数计算方式为：log2(p + 1) 向上取整。"),
     
    (r".*此时\s+\$s\[i\]\$.*可以表示为：前缀子问题最优值.*", 
     "• 此时的 s[i] 即为：前 i-k 个像素的最优存储位数，加上最后这一段的存储开销（k 乘以最大位宽，再加上 11 位段头）。"),
     
    (r"\$\$\s*s\[i\].*\\min_.*", 
     "递推公式：s[i] = 最小化 { s[i-k] + k × (该段最大位宽) } + 11，其中 k 的取值范围是 1 到 i 与 256 的最小值。"),
     
    (r".*边界条件：.*", 
     "• 边界条件：s[0] = 0。"),
     
    # Slide 5
    (r"构建一维\s*(?:DP|DP表|dp|dp表|动态规划)\s*表.*自底向上.*", 
     "构建一维状态数组 s[0..n]。自底向上从小到大计算，在计算每个位置时，通过遍历最后一段可能的所有长度 j，计算并填入最优值。"),
     
    (r".*外层循环\s*\$i\$.*依次求解.*", 
     "• 外层循环 i：遍历所有的前缀像素数，依次求解 s[1], s[2], ..., s[n]。"),
     
    (r".*内层循环\s*\$j\$.*遍历所有可能的.*", 
     "• 内层循环 j：遍历所有可能的最优分段长度 j（从 1 到最大限制），计算并取其最小值。"),
     
    (r"每次计算\s+\$s\[i\]\$.*值仅依赖于已算好.*", 
     "每次计算 s[i] 时，其值仅依赖于先前已经计算并保存好的较小子问题最优值 s[i-j]。"),
     
    # Slide 8
    (r"“?\s*\[!NOTE\]\s*”?|“?\s*\[!NOTE\]\s*”?", "设计提示："),
    
    (r"“?\s*设计提示：如果我们在实际.*", 
     "如果实际应用中只需要知道最小存储位数，无需输出具体的分段划分方案，可以直接计算到 s[n] 结束；若需给出具体划分边界，则必须执行第四步的回溯。"),
     
    (r"通过递归回溯已记录的每一段.*", 
     "通过递归回溯已经记录的每个位置的最优段长度，依次推导出每个分段的划分边界，最终输出具体方案。"),
     
    # Slide 10
    (r".*外层循环：执行.*", 
     "• 外层循环：执行 n 次。"),
     
    (r".*内层循环：最多迭代.*", 
     "• 内层循环：最多循环 256 次。因为最大段长是固定的常数限制，所以内层循环的执行时间为常数级别。"),
     
    (r".*总时间复杂度.*O\(n\).*", 
     "• 总时间复杂度：为 O(n)（线性时间复杂度），在大数据量下优势明显。"),
     
    (r".*使用了一维数组\s*s.*", 
     "• 空间复杂度：使用了一维数组记录状态，辅助空间与像素总数 n 成线性关系。"),
     
    (r".*总空间复杂度.*O\(n\).*", 
     "• 总空间复杂度：为 O(n)。"),
     
    # Slide 11
    (r"常数上限优化：通过对最大段长度.*成功优化至.*", 
     "常数上限优化：通过将最大分段长度限制在 256 像素（用8位表示），成功将原本是二次方复杂度的算法降到了线性复杂度。")
]

# LaTeX 公式通用清洗函数（用于未匹配到规则的行内公式）
def clean_latex_symbols(text):
    if not text:
        return text
    # 替换 \cdots 或者是 \dots 为 ...
    text = re.sub(r'\\(?:c)?dots', '...', text)
    # 替换 LaTeX 常见符号
    text = text.replace(r'\le', '<=')
    text = text.replace(r'\ge', '>=')
    text = text.replace(r'\in', '属于')
    text = text.replace(r'\cdot', '×')
    text = text.replace(r'\min', '最小值').replace(r'\max', '最大值')
    text = text.replace(r'\sum', '求和')
    text = re.sub(r'\\text\{\s*bits\s*\}', '位', text)
    text = re.sub(r'\\text\{\s*([^\}]+)\}', r'\1', text)
    text = text.replace(r'\{', '{').replace(r'\}', '}')
    text = text.replace(r'\left(', '(').replace(r'\right)', ')')
    text = text.replace(r'\left\{', '{').replace(r'\right\}', '}')
    text = text.replace(r'\lceil', '⌈').replace(r'\rceil', '⌉')
    # 去除 $ 符号
    text = text.replace('$', '')
    # 去除多余的反引号
    text = text.replace('`', '')
    return text

def clean_single_paragraph(text):
    if not text:
        return text
        
    # 1. 保护代码段落，不修改任何可能用作乘号或指针的单个星号 *
    if is_code_paragraph(text):
        # 仅去除 Markdown 的粗体加粗双星号
        return text.replace('**', '')
        
    # 2. 匹配规则库，转换成自然语言
    for pattern, replacement in RULES:
        if re.search(pattern, text):
            return replacement
            
    # 3. 通用 LaTeX 公式清洗
    if '$' in text:
        text = clean_latex_symbols(text)
        
    # 4. 移除 Markdown 遗留符号
    # 移除行首的标题标记 (例如 ###, ##, #)
    text = re.sub(r'^#+\s*', '', text)
    # 移除行首的无序列表标记 (例如 -, +, *)，后面带有空格
    text = re.sub(r'^[-\+\*]\s+', '', text)
    # 移除行首的有序列表标记 (例如 1., 2. 等)
    text = re.sub(r'^\d+[\.\)]\s+', '', text)
    # 移除行首的引用符号 >
    text = re.sub(r'^>\s*', '', text)
    
    # 5. 去除其它可能残留的 Markdown 标记
    text = text.replace('**', '')
    text = text.replace('*', '')
    text = text.replace('`', '')
    
    return text

def update_paragraph(paragraph, cleaned_text):
    full_text = paragraph.text
    if not full_text:
        return
        
    # 验证 cleaned_text 是否是 full_text 的子序列
    is_subseq = True
    idx = 0
    for char in cleaned_text:
        idx = full_text.find(char, idx)
        if idx == -1:
            is_subseq = False
            break
        idx += 1
        
    if is_subseq:
        # 使用双指针逐字符分配，保留格式
        cleaned_idx = 0
        for run in paragraph.runs:
            new_run_chars = []
            for char in run.text:
                if cleaned_idx < len(cleaned_text) and char == cleaned_text[cleaned_idx]:
                    new_run_chars.append(char)
                    cleaned_idx += 1
            run.text = "".join(new_run_chars)
    else:
        # 如果由于重写段落（不是子序列），直接重置整个段落的 text
        paragraph.text = cleaned_text

def process_text_frame(tf):
    for paragraph in tf.paragraphs:
        cleaned = clean_single_paragraph(paragraph.text)
        update_paragraph(paragraph, cleaned)

def process_shapes(shapes):
    for shape in shapes:
        # 处理普通文本框
        if shape.has_text_frame:
            process_text_frame(shape.text_frame)
        
        # 处理表格
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    process_text_frame(cell.text_frame)
                    
        # 处理组合形状 (GroupShape, shape_type == 6)
        if shape.shape_type == 6:
            try:
                process_shapes(shape.shapes)
            except Exception as e:
                print(f"Warning processing sub-shapes: {e}")

def main():
    # 从最原始的文件重新读取
    path = r"C:\Users\Dell\Downloads\image_compression_presentation_editable.pptx"
    out_path = r"C:\Users\Dell\Downloads\图像压缩.pptx"
    
    if not os.path.exists(path):
        print(f"Error: File not found at {path}")
        return
        
    prs = Presentation(path)
    
    # 遍历所有幻灯片
    for i, slide in enumerate(prs.slides):
        process_shapes(slide.shapes)
        # 处理备注
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            process_text_frame(slide.notes_slide.notes_text_frame)
                                
    prs.save(out_path)
    print(f"Successfully cleaned PPT and saved to: {out_path}")

if __name__ == '__main__':
    main()
