import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# 移动幻灯片位置
def move_slide(prs, slide, new_index):
    sldIdLst = prs.slides._sldIdLst
    slide_id = prs.slides.index(slide)
    slide_id_element = sldIdLst[slide_id]
    sldIdLst.remove(slide_id_element)
    sldIdLst.insert(new_index, slide_id_element)

# 在卡片内添加多级列表格式文本
def add_card(slide, left, top, width, height, title_text, body_text):
    # 1. 绘制背景圆角矩形
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(245, 247, 250)  # 淡淡的蓝灰底色
    shape.line.color.rgb = RGBColor(218, 223, 230)       # 极其微弱的边框线
    shape.line.width = Pt(1.5)
    
    # 2. 添加卡片内部标题文本框
    title_box = slide.shapes.add_textbox(left + Pt(10), top + Pt(12), width - Pt(20), Pt(35))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_top = 0
    tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.bold = True
    p_title.font.size = Pt(14)
    p_title.font.color.rgb = RGBColor(15, 23, 42)  # 深石板色
    p_title.font.name = "微软雅黑"
    
    # 3. 添加卡片正文文本框
    body_box = slide.shapes.add_textbox(left + Pt(8), top + Pt(46), width - Pt(16), height - Pt(55))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    tf_body.margin_top = 0
    tf_body.margin_bottom = 0
    
    lines = body_text.split('\n')
    is_first = True
    for line in lines:
        if not line.strip():
            continue
            
        if is_first:
            p = tf_body.paragraphs[0]
            is_first = False
        else:
            p = tf_body.add_paragraph()
        
        p.font.name = "微软雅黑"
        stripped = line.strip()
        
        # 判断多级项目列表符号并排版
        if stripped.startswith("•"):
            p.text = stripped
            p.font.bold = True
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor(30, 41, 59) # 深蓝黑色
            p.space_before = Pt(4)
            p.space_after = Pt(2)
        elif stripped.startswith("-"):
            # 缩进子项目，使用 4 个半角空格进行模拟排版
            p.text = "    " + stripped
            p.font.bold = False
            p.font.size = Pt(9)
            p.font.color.rgb = RGBColor(71, 85, 105) # 蓝灰色
            p.space_after = Pt(2)
        else:
            p.text = stripped
            p.font.bold = False
            p.font.size = Pt(9)
            p.font.color.rgb = RGBColor(100, 116, 139) # 灰色
            p.space_after = Pt(2)

def main():
    # 直接在已经被 clean_ppt.py 清洗过的图像压缩.pptx 上添加页面
    output_path = r"C:\Users\Dell\Downloads\图像压缩.pptx"
    
    if not os.path.exists(output_path):
        print(f"Error: Base file not found at {output_path}")
        return
        
    prs = Presentation(output_path)
    
    try:
        slide_layout = prs.slide_layouts[5]
    except IndexError:
        slide_layout = prs.slide_layouts[0]
        
    W = prs.slide_width
    H = prs.slide_height
    
    # ------------------ 新增幻灯片 A (i=1..3) ------------------
    slide_a = prs.slides.add_slide(slide_layout)
    slide_a.shapes.title.text = "手算递推实例（一）：计算 i = 1 到 3"
    
    card_width = W * 0.28
    card_height = H * 0.64
    card_top = H * 0.23
    spacing = W * 0.04
    left_margin = W * 0.04
    
    # 卡片 1 (i = 1)
    body_i1 = (
        "• 处理像素：p[1] = 10（初始位宽 4位）\n"
        "• 默认划分（最后一段长度为 1）：\n"
        "  - 划分情况：║ [ 10 ]\n"
        "  - 空间开销：s[0] + 1 * 4 = 0 + 4 = 4位\n"
        "• 决策与结算：\n"
        "  - 唯一方案，直接采纳\n"
        "  - 加上段头开销 11位：s[1] = 4 + 11 = 15位\n"
        "  - 记录最优最后段长 l[1] = 1，位深 b[1] = 4"
    )
    add_card(slide_a, left_margin, card_top, card_width, card_height, "第 1 步：i = 1", body_i1)
    
    # 卡片 2 (i = 2)
    body_i2 = (
        "• 处理像素：p[2] = 12（初始位宽 4位）\n"
        "• 试探 j = 1（第2个像素单独划为一段）：\n"
        "  - 划分情况：[ 10 ] ║ [ 12 ]\n"
        "  - 空间开销：s[1] + 1 * 4 = 15 + 4 = 19位\n"
        "• 试探 j = 2（前2个像素合并为一段）：\n"
        "  - 划分情况：║ [ 10, 12 ]\n"
        "  - 空间开销：s[0] + 2 * 4 = 0 + 8 = 8位\n"
        "• 决策与结算：\n"
        "  - 因 8 < 19，选择合并方案 (j=2)\n"
        "  - 加上段头开销 11位：s[2] = 8 + 11 = 19位\n"
        "  - 记录最优最后段长 l[2] = 2，位深 b[2] = 4"
    )
    add_card(slide_a, left_margin + card_width + spacing, card_top, card_width, card_height, "第 2 步：i = 2", body_i2)
    
    # 卡片 3 (i = 3)
    body_i3 = (
        "• 处理像素：p[3] = 15（初始位宽 4位）\n"
        "• 试探 j = 1（第3个像素单独划为一段）：\n"
        "  - 划分情况：[ 10, 12 ] ║ [ 15 ]\n"
        "  - 空间开销：s[2] + 1 * 4 = 19 + 4 = 23位\n"
        "• 试探 j = 2（最后两个像素合并为一段）：\n"
        "  - 划分情况：[ 10 ] ║ [ 12, 15 ]\n"
        "  - 空间开销：s[1] + 2 * 4 = 15 + 8 = 23位\n"
        "• 试探 j = 3（前3个像素全部合并）：\n"
        "  - 划分情况：║ [ 10, 12, 15 ]\n"
        "  - 空间开销：s[0] + 3 * 4 = 0 + 12 = 12位\n"
        "• 决策与结算：\n"
        "  - 因 12 < 23，选择全合并方案 (j=3)\n"
        "  - 加上段头开销 11位：s[3] = 12 + 11 = 23位\n"
        "  - 记录最优最后段长 l[3] = 3，位深 b[3] = 4"
    )
    add_card(slide_a, left_margin + (card_width + spacing) * 2, card_top, card_width, card_height, "第 3 步：i = 3", body_i3)
    
    # ------------------ 新增幻灯片 B (i=4..5) ------------------
    slide_b = prs.slides.add_slide(slide_layout)
    slide_b.shapes.title.text = "手算递推实例（二）：关键决策 i = 4 与 i = 5"
    
    # 卡片 1 (i = 4)
    body_i4 = (
        "• 处理像素：p[4] = 255（初始位宽 8位）\n"
        "• 试探 j = 1（大数 255 单独划为一段）：\n"
        "  - 划分情况：[ 10, 12, 15 ] ║ [ 255 ]\n"
        "  - 空间开销：s[3] + 1 * 8 = 23 + 8 = 31位\n"
        "• 试探 j = 2 到 4（尝试将大数并入）：\n"
        "  - j = 2：s[2] + 2 * 8 = 19 + 16 = 35位\n"
        "  - j = 3：s[1] + 3 * 8 = 15 + 24 = 39位\n"
        "  - j = 4：s[0] + 4 * 8 = 0 + 32 = 32位\n"
        "• 决策与结算：\n"
        "  - 因 31 为各方案最小值，选择 j=1 (单飞)\n"
        "  - 加上段头开销 11位：s[4] = 31 + 11 = 42位\n"
        "  - 记录最优最后段长 l[4] = 1，位深 b[4] = 8"
    )
    add_card(slide_b, left_margin, card_top, card_width, card_height, "第 4 步：i = 4 (大数单飞)", body_i4)
    
    # 卡片 2 (i = 5)
    body_i5 = (
        "• 处理像素：p[5] = 1（初始位宽 1位）\n"
        "• 试探 j = 1（第5个像素单独划为一段）：\n"
        "  - 划分情况：[10,12,15] [255] ║ [1]\n"
        "  - 空间开销：s[4] + 1 * 1 = 42 + 1 = 43位\n"
        "• 试探 j = 2（大数 255 与 1 合并）：\n"
        "  - 划分情况：[ 10, 12, 15 ] ║ [ 255, 1 ]\n"
        "  - 空间开销：s[3] + 2 * 8 = 23 + 16 = 39位\n"
        "• 试探 j = 3 到 5：\n"
        "  - j = 3：s[2] + 3 * 8 = 19 + 24 = 43位\n"
        "  - j = 5：s[0] + 5 * 8 = 0 + 40 = 40位\n"
        "• 决策与结算：\n"
        "  - 因 39 最优，选择合并方案 (j=2)\n"
        "  - 加上段头开销 11位：s[5] = 39 + 11 = 50位\n"
        "  - 记录最优最后段长 l[5] = 2，位深 b[5] = 8"
    )
    add_card(slide_b, left_margin + card_width + spacing, card_top, card_width, card_height, "第 5 步：i = 5 (重新合并)", body_i5)
    
    # 卡片 3 (总结)
    body_sum = (
        "• 最终最优分段方案：\n"
        "  - 根据前 5 个像素回溯，最终确定划分：\n"
        "• 第 1 分段 [像素 1 到 3]：\n"
        "  - 包含数值：[ 10, 12, 15 ]（位宽 4位）\n"
        "  - 存储空间：3 * 4 + 11 = 23位\n"
        "• 第 2 分段 [像素 4 到 5]：\n"
        "  - 包含数值：[ 255, 1 ]（位宽 8位）\n"
        "  - 存储空间：2 * 8 + 11 = 27位\n"
        "• 总存储开销：\n"
        "  - 两段总空间合计为：23 + 27 = 50位\n"
        "  - (因为数据量小导致头部开销占比大；在大规模像素下，变长压缩优势将极其明显)"
    )
    add_card(slide_b, left_margin + (card_width + spacing) * 2, card_top, card_width, card_height, "前 5 像素最优解总结", body_sum)
    
    # ------------------ 调整位次 ------------------
    # slide_a 移到索引 6，slide_b 移到索引 7
    move_slide(prs, slide_a, 6)
    move_slide(prs, slide_b, 7)
    
    prs.save(output_path)
    print("Successfully generated clean PPT without corruptions.")

if __name__ == '__main__':
    main()
