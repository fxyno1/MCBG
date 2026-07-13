import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def update_ppt(file_path):
    backup_path = file_path.replace(".pptx", "_backup_before_table.pptx")
    if os.path.exists(file_path):
        shutil.copy(file_path, backup_path)
    
    try:
        prs = Presentation(file_path)
        
        # 使用空白布局新增一页幻灯片 (通常布局6是空白页)
        try:
            slide_layout = prs.slide_layouts[6] 
        except IndexError:
            slide_layout = prs.slide_layouts[0]
            
        slide = prs.slides.add_slide(slide_layout)
        
        # 添加标题
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.8))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "图像压缩算法计算过程演示 (i = 1 到 5)"
        p.font.size = Pt(28)
        p.font.bold = True
        
        # 添加文字说明
        expBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(9), Inches(1))
        exp_tf = expBox.text_frame
        exp_tf.word_wrap = True
        exp_p = exp_tf.paragraphs[0]
        exp_p.text = "在计算前i个像素的最优分段空间s[i]时，程序通过循环尝试将最后j个像素划分为新段。当前开销 = s[i-j] + j × bmax。我们选出开销最小的方案，加上固定的11位段头开销，即得到最优决策s[i]。"
        exp_p.font.size = Pt(14)
        
        # 添加表格 6行 5列
        rows, cols = 6, 5
        left, top, width, height = Inches(0.2), Inches(1.6), Inches(9.6), Inches(4.5)
        table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        table = table_shape.table
        
        # 设置列宽
        table.columns[0].width = Inches(0.8)
        table.columns[1].width = Inches(1.3)
        table.columns[2].width = Inches(3.8)
        table.columns[3].width = Inches(2.2)
        table.columns[4].width = Inches(1.5)
        
        # 表头数据
        headers = ["阶段(i)", "当前像素", "各可能段长(j)的计算公式：s[i-j] + j * bmax", "循环比较数组", "最终最优决策"]
        for col_idx, header in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = header
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.bold = True
                paragraph.alignment = PP_ALIGN.CENTER
        
        # 表格内容
        data = [
            ["i = 1", "p[1]=10\n(位宽4)", "j=1: s[0] + 1×4 = 0 + 4 = 4", "[ 4 ]", "选最小值4\ns[1] = 15位"],
            ["i = 2", "p[2]=12\n(位宽4)", "j=1: s[1] + 1×4 = 15 + 4 = 19\nj=2: s[0] + 2×4 = 0 + 8 = 8", "[ 19, 8 ]", "选合并方案(8)\ns[2] = 19位"],
            ["i = 3", "p[3]=15\n(位宽4)", "j=1: s[2] + 1×4 = 19 + 4 = 23\nj=2: s[1] + 2×4 = 15 + 8 = 23\nj=3: s[0] + 3×4 = 0 + 12 = 12", "[ 23, 23, 12 ]", "选全合并(12)\ns[3] = 23位"],
            ["i = 4", "p[4]=255\n(位宽8)", "j=1: s[3] + 1×8 = 23 + 8 = 31\nj=2: s[2] + 2×8 = 19 + 16 = 35\nj=3: s[1] + 3×8 = 15 + 24 = 39\nj=4: s[0] + 4×8 = 0 + 32 = 32", "[ 31, 35, 39, 32 ]", "选单飞(31)\ns[4] = 42位"],
            ["i = 5", "p[5]=1\n(位宽1)", "j=1: s[4] + 1×1 = 42 + 1 = 43\nj=2: s[3] + 2×8 = 23 + 16 = 39\nj=3: s[2] + 3×8 = 19 + 24 = 43\nj=4: s[1] + 4×8 = 15 + 32 = 47\nj=5: s[0] + 5×8 = 0 + 40 = 40", "[ 43, 39, 43, 47, 40 ]", "选合并(39)\ns[5] = 50位"]
        ]
        
        for row_idx, row_data in enumerate(data):
            for col_idx, cell_data in enumerate(row_data):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = cell_data
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(13)
        
        prs.save(file_path)
        print("Success: Appended table slide to PPTX")
    except Exception as e:
        print(f"Failed to modify PPTX: {e}")

if __name__ == "__main__":
    ppt_path = r"C:\Users\Dell\Downloads\图像压缩.pptx"
    update_ppt(ppt_path)
